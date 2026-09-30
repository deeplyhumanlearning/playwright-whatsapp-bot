import json
import random
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
from playwright.sync_api import sync_playwright

BASE_DIR = Path(__file__).resolve().parent
CONTACTS_FILE = BASE_DIR / "contacts.xlsx"
OUTPUT_DIR = BASE_DIR / "output"
SCREENSHOT_DIR = OUTPUT_DIR / "screenshots"
SESSION_DIR = BASE_DIR / "whatsapp_session"

OUTPUT_DIR.mkdir(exist_ok=True)
SCREENSHOT_DIR.mkdir(exist_ok=True)


def pause(page, minimum=2, maximum=5):
    page.wait_for_timeout(random.randint(minimum * 1000, maximum * 1000))


def clean_phone(phone):
    return re.sub(r"\D", "", str(phone))


def read_contacts():
    if not CONTACTS_FILE.exists():
        raise FileNotFoundError(f"Missing {CONTACTS_FILE.name}")

    workbook = load_workbook(CONTACTS_FILE, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    if not rows:
        return []

    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    required = {"Name", "Phone", "Message"}
    if not required.issubset(headers):
        raise ValueError("contacts.xlsx must contain Name, Phone and Message columns")

    indexes = {header: headers.index(header) for header in required}
    contacts = []

    for row in rows[1:]:
        name = row[indexes["Name"]] if len(row) > indexes["Name"] else None
        phone = row[indexes["Phone"]] if len(row) > indexes["Phone"] else None
        message = row[indexes["Message"]] if len(row) > indexes["Message"] else None

        if not name or not phone:
            continue

        contacts.append(
            {
                "name": str(name).strip(),
                "phone": str(phone).strip(),
                "message": str(message).strip() if message else "",
            }
        )

    return contacts


def find_search_box(page):
    selectors = [
        'div[contenteditable="true"][aria-label*="Search"]',
        'div[contenteditable="true"][data-tab="3"]',
    ]
    for selector in selectors:
        locator = page.locator(selector).first
        if locator.count():
            return locator
    raise PlaywrightTimeoutError("WhatsApp search box was not found")


def open_contact(page, name, phone):
    search_box = find_search_box(page)

    for search_term in [name, phone]:
        if not search_term:
            continue

        search_box.click()
        search_box.fill(search_term)
        pause(page)

        try:
            page.wait_for_selector('div[role="listitem"]', timeout=10000)
        except PlaywrightTimeoutError:
            continue

        candidates = page.locator('div[role="listitem"]').filter(has_text=search_term)
        if candidates.count() == 0:
            continue

        try:
            candidates.first.click()
            pause(page)
            return True
        except PlaywrightTimeoutError:
            continue

    return False


def find_message_box(page):
    selectors = [
        'div[contenteditable="true"][aria-label="Type a message"]',
        'div[contenteditable="true"][data-tab="10"]',
    ]
    for selector in selectors:
        locator = page.locator(selector).last
        if locator.count():
            return locator
    raise PlaywrightTimeoutError("Message box was not found")


def send_message(page, message):
    message_box = find_message_box(page)
    message_box.click()
    message_box.fill(message)
    pause(page)
    message_box.press("Enter")
    pause(page)

    sent = page.locator("div.message-out").filter(has_text=message).last
    sent.wait_for(state="visible", timeout=10000)
    return sent


def extract_messages(page):
    messages = page.locator("div.message-in")
    count = messages.count()
    extracted = []

    for index in range(max(0, count - 3), count):
        text = messages.nth(index).inner_text().strip()
        if text:
            extracted.append({"sender": "contact", "message": text})

    return extracted


def save_report(results, run_date):
    json_file = OUTPUT_DIR / f"whatsapp_report_{run_date}.json"
    xlsx_file = OUTPUT_DIR / f"whatsapp_report_{run_date}.xlsx"

    with json_file.open("w", encoding="utf-8") as file:
        json.dump(
            {
                "run_date": run_date,
                "contacts": results,
            },
            file,
            indent=2,
            ensure_ascii=False,
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.append(["Name", "Phone", "Status", "Message", "Last 3 Messages", "Error"])

    for result in results:
        last_messages = " | ".join(
            f"{item['sender']}: {item['message']}" for item in result["last_3_messages"]
        )
        sheet.append(
            [
                result["name"],
                result["phone"],
                result["status"],
                result["message"],
                last_messages,
                result["error"],
            ]
        )

    workbook.save(xlsx_file)
    return json_file, xlsx_file


def main():
    contacts = read_contacts()
    if not contacts:
        print("No contacts found in contacts.xlsx")
        return

    run_date = datetime.now().strftime("%Y-%m-%d")
    results = []

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch_persistent_context(
            str(SESSION_DIR),
            headless=False,
            viewport={"width": 1280, "height": 900},
        )
        page = browser.pages[0] if browser.pages else browser.new_page()
        page.goto("https://web.whatsapp.com", wait_until="domcontentloaded")

        print("Waiting for WhatsApp Web. Scan the QR code on the first run.")
        try:
            page.wait_for_selector(
                'div[contenteditable="true"][aria-label*="Search"], div[contenteditable="true"][data-tab="3"]',
                timeout=120000,
            )
        except PlaywrightTimeoutError:
            browser.close()
            raise RuntimeError("WhatsApp Web login timed out")

        for contact in contacts:
            result = {
                "name": contact["name"],
                "phone": contact["phone"],
                "message": "",
                "status": "failed",
                "last_3_messages": [],
                "screenshot": "",
                "error": "",
            }

            try:
                if not open_contact(page, contact["name"], contact["phone"]):
                    raise ValueError("Contact not found")

                message = contact["message"] or ""
                message = message.replace("{name}", contact["name"])
                if not message:
                    raise ValueError("Message is empty")

                result["message"] = message
                sent_message = send_message(page, message)

                screenshot_name = (
                    f"{datetime.now().strftime('%H%M%S')}_{clean_phone(contact['phone'])}.png"
                )
                screenshot_path = SCREENSHOT_DIR / screenshot_name
                sent_message.screenshot(path=str(screenshot_path))
                result["screenshot"] = str(screenshot_path.relative_to(BASE_DIR))
                result["status"] = "sent"
                result["last_3_messages"] = extract_messages(page)

            except Exception as error:
                result["error"] = str(error)

            results.append(result)
            pause(page)

        json_file, xlsx_file = save_report(results, run_date)
        print(f"JSON report: {json_file}")
        print(f"Excel report: {xlsx_file}")
        browser.close()


if __name__ == "__main__":
    main()
